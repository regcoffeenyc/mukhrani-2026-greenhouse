#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Builds the strip-foundation take-off — excavation, concrete, cement, sand,
gravel, reinforcement and formwork for a perimeter strip footing.

This is a QUANTITY calculator, not a design. The section (width x depth) and the
reinforcement come from whoever engineers the foundation; type them in and this
tells you what to order. Every figure is a live formula off the yellow inputs.

Rebar mass uses the standard d^2/162 kg per metre, so any bar diameter works.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = str(Path(__file__).with_name("Foundation_Calculator.xlsx"))

F = "Arial"
TITLE = Font(name=F, size=15, bold=True, color="1F3864")
H1 = Font(name=F, size=10, bold=True, color="FFFFFF")
BOLD = Font(name=F, size=10, bold=True)
TXT = Font(name=F, size=10)
NOTE = Font(name=F, size=9, italic=True, color="666666")
WARN = Font(name=F, size=10, bold=True, color="9C0006")
BIG = Font(name=F, size=12, bold=True, color="1F3864")

FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
FILL_WARN = PatternFill("solid", fgColor="FFC7CE")

thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

N0, N2, N3 = '#,##0', '#,##0.00', '#,##0.000'

wb = Workbook()
ws = wb.active
ws.title = "Foundation"
ws.sheet_view.showGridLines = False
for col, w in {"A": 48, "B": 14, "C": 8, "D": 54}.items():
    ws.column_dimensions[col].width = w


def put(cell, value, font=TXT, fmt=None, fill=None, border=None, align=None):
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return c


def section(row, label):
    ws.merge_cells(f"A{row}:D{row}")
    put(f"A{row}", label, H1, fill=FILL_H)
    ws.row_dimensions[row].height = 20


def inp(row, label, value, fmt, note=""):
    put(f"A{row}", label, TXT, border=BOX)
    put(f"B{row}", value, BOLD, fmt=fmt, fill=FILL_IN, border=BOX, align="right")
    put(f"D{row}", note, NOTE)


def out(row, label, formula, fmt, note="", big=False):
    put(f"A{row}", label, BOLD if big else TXT, border=BOX,
        fill=FILL_OUT if big else None)
    put(f"B{row}", formula, BIG if big else BOLD, fmt=fmt,
        fill=FILL_OUT if big else None, border=BOX, align="right")
    put(f"D{row}", note, NOTE)


put("A1", "საძირკვლის კალკულატორი — ლენტური / STRIP FOUNDATION CALCULATOR", TITLE)
put("A2", "ეს რაოდენობის კალკულატორია, არა პროექტი. კვეთა და არმატურა ინჟინრისგან/ჯიესენისგან უნდა მოვიდეს. / "
          "This counts materials. It does not design anything — the section and the rebar must come from "
          "the people engineering the foundation.", NOTE)

ws.merge_cells("A3:D3")
put("A3", "⚠ სანამ შეუკვეთავთ: აიღეთ კვეთა და არმატურა ჯიესენ გრუპის ნახაზებიდან. ქვემოთ ჩაწერილი ციფრები "
          "მსუბუქი პერიმეტრის კედლის ტიპიური მნიშვნელობებია, თქვენი ობიექტის პროექტი არ არის. / "
          "⚠ Before ordering: take the section and the reinforcement from GSN Group's drawings. The numbers "
          "below are typical for a light perimeter wall — they are not a design for your site.",
    WARN, fill=FILL_WARN, align="left")
ws.row_dimensions[3].height = 34

# ---------------- geometry ----------------
section(5, "1. გეომეტრია / GEOMETRY")
inp(6, "სიგრძე — პერიმეტრი / Length (perimeter)", 180, N2, "მეტრი / metres")
inp(7, "საძირკვლის სიგანე / Footing width", 0.40, N3,
    "მეტრი / metres — from the drawings")
inp(8, "საძირკვლის სიმაღლე — ბეტონი / Footing depth (concrete)", 0.50, N3,
    "მეტრი / metres — from the drawings")
inp(9, "ხრეშის ბალიში / Compacted gravel bed", 0.10, N3,
    "მეტრი / metres under the footing. 0 if none")
inp(10, "მოსასწორებელი ბეტონი / Blinding (lean concrete)", 0.05, N3,
    "მეტრი / metres. 0 if none")
inp(11, "თხრილის დამატებითი სიგანე / Extra trench width for working space",
    0.20, N3, "მეტრი, ორივე მხარეს ჯამში / metres total, both sides")

# ---------------- concrete ----------------
section(13, "2. ბეტონის შემადგენლობა / CONCRETE MIX (per m³)")
inp(14, "ცემენტი 1 მ³-ზე — კონსტრუქციული / Cement per m³ — structural", 320, N0,
    "კგ. C20/25 (M250) ≈ 320 / kg")
inp(15, "ქვიშა 1 მ³-ზე / Sand per m³", 0.45, N2, "მ³ / m³")
inp(16, "ღორღი 1 მ³-ზე / Crushed stone per m³", 0.85, N2, "მ³ / m³")
inp(17, "წყალი 1 მ³-ზე / Water per m³", 180, N0, "ლიტრი / litres")
inp(18, "ცემენტი 1 მ³-ზე — მოსასწორებელი / Cement per m³ — blinding", 200, N0,
    "კგ, მჭლე ნარევი / kg, lean mix")
inp(19, "ბეტონის მარაგი / Concrete waste allowance", 0.05, '0%', "")
inp(20, "ცემენტის ტომარა / Cement bag weight", 50, N0, "კგ / kg")

# ---------------- rebar ----------------
section(22, "3. არმატურა / REINFORCEMENT")
inp(23, "გრძივი ღერო — რაოდენობა / Longitudinal bars — how many", 4, N0,
    "ნახაზიდან / from the drawings")
inp(24, "გრძივი ღერო — დიამეტრი / Longitudinal bar diameter", 12, N0, "მმ / mm")
inp(25, "ხომუთი — დიამეტრი / Stirrup diameter", 8, N0, "მმ / mm")
inp(26, "ხომუთის ბიჯი / Stirrup spacing", 0.30, N3, "მეტრი / metres")
inp(27, "დამცავი ფენა / Concrete cover", 0.04, N3, "მეტრი, ყოველ მხარეს / metres each face")
inp(28, "გადაბმა და კუთხეები / Laps and corners allowance", 0.10, '0%',
    "გრძივ ღეროებზე / on the longitudinal steel")
inp(29, "არმატურის მარაგი / Steel waste allowance", 0.05, '0%', "")
inp(30, "შესაკრავი მავთული / Binding wire", 10, N0, "კგ 1 ტონა არმატურაზე / kg per tonne")

# ---------------- formwork ----------------
section(32, "4. ყალიბი / FORMWORK")
inp(33, "ყალიბის სიმაღლე / Formwork height", 0.00, N3,
    "მეტრი. 0 = თხრილში ჩაისხმება, ყალიბი არ სჭირდება / "
    "metres. 0 means poured against the trench — no formwork")

# ---------------- working ----------------
section(35, "5. გაანგარიშება / THE WORKING")
out(36, "თხრილის კვეთა / Trench cross-section", "=(B7+B11)*(B8+B9+B10)", N3, "მ² / m²")
out(37, "გრუნტის მოცულობა / Excavation", "=B6*B36", N2, "მ³ / m³")
out(38, "1 მ-ზე ბეტონი / Concrete per metre of footing", "=B7*B8", N3, "მ³ / m³")
out(39, "ხომუთის რაოდენობა / Stirrup count", "=ROUNDUP(B6/B26,0)+1", N0, "ცალი / pieces")
out(40, "ერთი ხომუთის სიგრძე / Length of one stirrup",
    "=2*(B7-2*B27)+2*(B8-2*B27)+0.10", N3, "მეტრი, კაუჭებით / metres, hooks included")
out(41, "არმატურის წონა 1 მ-ზე — გრძივი / kg per m — longitudinal", "=B24^2/162", N3,
    "d²/162 — the standard formula")
out(42, "არმატურის წონა 1 მ-ზე — ხომუთი / kg per m — stirrup", "=B25^2/162", N3, "")

# ---------------- order ----------------
section(44, "6. საყიდელი / WHAT TO ORDER")
out(45, "ხრეში — ბალიში / Gravel for the bed", "=ROUNDUP(B6*B7*B9*10,0)/10", N2, "მ³ / m³")
out(46, "მოსასწორებელი ბეტონი / Blinding concrete", "=ROUNDUP(B6*B7*B10*10,0)/10", N2, "მ³ / m³")
out(47, "კონსტრუქციული ბეტონი / STRUCTURAL CONCRETE",
    "=ROUNDUP(B6*B38*(1+B19)*10,0)/10", N2, "მ³ / m³, მარაგით / waste included", big=True)
out(48, "ცემენტი — ჯამში / CEMENT, total",
    "=ROUNDUP((B47*B14+B46*B18)/B20,0)", N0, "ტომარა 50 კგ / bags", big=True)
out(49, "ქვიშა / SAND", "=ROUNDUP(B47*B15*10,0)/10", N2, "მ³ / m³", big=True)
out(50, "ღორღი / CRUSHED STONE", "=ROUNDUP(B47*B16*10,0)/10", N2, "მ³ / m³", big=True)
out(51, "წყალი / Water", "=ROUNDUP(B47*B17,0)", N0, "ლიტრი / litres")
out(52, "არმატურა — გრძივი / Steel — longitudinal",
    "=ROUNDUP(B6*B23*(1+B28)*B41*(1+B29),0)", N0, "კგ / kg", big=True)
out(53, "არმატურა — ხომუთი / Steel — stirrups",
    "=ROUNDUP(B39*B40*B42*(1+B29),0)", N0, "კგ / kg", big=True)
out(54, "არმატურა — სულ / STEEL, total", "=B52+B53", N0, "კგ / kg", big=True)
out(55, "შესაკრავი მავთული / Binding wire", "=ROUNDUP(B54/1000*B30,0)", N0, "კგ / kg")
out(56, "ყალიბი / Formwork", "=ROUNDUP(B6*2*B33*10,0)/10", N2, "მ² / m²")

# ---------------- notes ----------------
section(58, "7. წასაკითხი / READ THIS BEFORE ORDERING")
notes = [
    "კვეთა და არმატურა აქ ტიპიურია მსუბუქი პერიმეტრის კედლისთვის — 0.80 მ ბლოკი. ეს არ არის თქვენი "
    "ობიექტის პროექტი. აიღეთ ნახაზი ჯიესენ გრუპისგან და ჩაწერეთ მათი ციფრები. / "
    "The section and rebar here are typical for a light 0.80 m perimeter wall. They are not a design for "
    "your site. Get the drawing from GSN Group and type their numbers in.",

    "გაყინვის სიღრმე: აღმოსავლეთ საქართველოში საძირკველი გაყინვის ზონის ქვემოთ უნდა ჩავიდეს. თუ სათბურის "
    "კარკასიც ამ საძირკველზე დგება, დატვირთვა სხვაა და ინჟინერი სჭირდება. / "
    "Frost depth: the footing must sit below it. And if the greenhouse frame lands on this same foundation, "
    "the loads are different and an engineer has to size it.",

    "38 მ³ ბეტონი ხელით ურევად ბევრია. მიწოდებული ბეტონი (მიქსერით) ჩვეულებრივ იაფიცაა და უკეთესიც — "
    "აიღეთ ფასი ორივეზე. / "
    "38 m³ is a lot of concrete to mix by hand. Ready-mix delivered is usually both cheaper and better at "
    "this volume — price both.",

    "ეს რაოდენობაა, არა ფასი. ფასი ხარჯების ლეჯერში შედის ყიდვის შემდეგ. / "
    "These are quantities, not costs. The cost goes into the expense ledger when it is bought.",
]
for i, n in enumerate(notes):
    r = 59 + i
    ws.merge_cells(f"A{r}:D{r}")
    put(f"A{r}", "· " + n, NOTE, align="left")
    ws.row_dimensions[r].height = 30

wb.save(OUT)
print("saved", OUT)
