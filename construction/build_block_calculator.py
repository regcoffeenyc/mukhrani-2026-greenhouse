#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Builds the perimeter blockwork calculator — blocks, mortar, cement and sand.

Everything is a live formula off the yellow input cells, so changing the
perimeter, the height or the block size re-computes the whole sheet.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = str(Path(__file__).with_name("Block_Calculator.xlsx"))

F = "Arial"
TITLE = Font(name=F, size=15, bold=True, color="1F3864")
H1 = Font(name=F, size=10, bold=True, color="FFFFFF")
BOLD = Font(name=F, size=10, bold=True)
TXT = Font(name=F, size=10)
NOTE = Font(name=F, size=9, italic=True, color="666666")
BIG = Font(name=F, size=12, bold=True, color="1F3864")

FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")   # yellow = you may edit
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")  # green  = the answer
FILL_SEC = PatternFill("solid", fgColor="D9E2F3")

thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

N2 = '#,##0.00'
N3 = '#,##0.000'
N0 = '#,##0'

wb = Workbook()
ws = wb.active
ws.title = "Calculator"
ws.sheet_view.showGridLines = False
for col, w in {"A": 46, "B": 14, "C": 10, "D": 52}.items():
    ws.column_dimensions[col].width = w


def put(cell, value, font=TXT, fmt=None, fill=None, border=None,
        align=None, note=None):
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if note:
        c.comment = Comment(note, "Model")
    return c


def section(row, label):
    ws.merge_cells(f"A{row}:D{row}")
    put(f"A{row}", label, H1, fill=FILL_H)
    ws.row_dimensions[row].height = 20


def inp(row, label, value, fmt, note=""):
    put(f"A{row}", label, TXT, border=BOX)
    put(f"B{row}", value, BOLD, fmt=fmt, fill=FILL_IN, border=BOX, align="right")
    put(f"D{row}", note, NOTE)


def out(row, label, formula, fmt, note="", big=False):
    put(f"A{row}", label, BOLD if big else TXT, border=BOX,
        fill=FILL_OUT if big else None)
    put(f"B{row}", formula, BIG if big else BOLD, fmt=fmt,
        fill=FILL_OUT if big else None, border=BOX, align="right")
    put(f"D{row}", note, NOTE)


put("A1", "ბლოკის კალკულატორი — პერიმეტრის კედელი / BLOCKWORK CALCULATOR — perimeter wall", TITLE)
put("A2", "ყვითელი უჯრები შეცვალეთ. მწვანე — შედეგი. ყველაფერი ფორმულაა. / "
          "Edit the yellow cells. Green is the answer. Everything recalculates.", NOTE)

# ---------------- inputs ----------------
section(4, "1. რას ვაშენებთ / WHAT WE ARE BUILDING")
inp(5, "პერიმეტრი / Perimeter", 180, N2,
    "მეტრი / metres — the full run of wall")
inp(6, "სიმაღლე / Wall height", 0.80, N2,
    "მეტრი / metres. 0.80 m = 80 სმ")
inp(7, "ღიობები — გამოსაკლები ფართობი / Openings to deduct", 0, N2,
    "მ² / m². ჭიშკარი, კარი. 0 = ღიობის გარეშე / gates and doors; 0 means none")

section(9, "2. ბლოკი და ნაკერი / BLOCK AND JOINT")
inp(10, "ბლოკის სიგრძე / Block length", 0.40, N3, "მეტრი / metres")
inp(11, "ბლოკის სიმაღლე / Block height", 0.20, N3, "მეტრი / metres")
inp(12, "ბლოკის სისქე — კედლის სისქე / Block width = wall thickness", 0.20, N3,
    "მეტრი / metres. ბრტყლად დაწყობილი / laid flat")
inp(13, "ნაკერის სისქე / Mortar joint", 0.01, N3,
    "მეტრი / metres. 0.01 m = 1 სმ, ჩვეულებრივი / 1 cm is normal")
inp(14, "ბლოკის მარაგი / Block waste allowance", 0.05, '0%',
    "გატეხილი და მოჭრილი / breakage and cuts. 5% is normal")

section(16, "3. ხსნარი / MORTAR")
inp(17, "ცემენტი : ქვიშა / Cement : sand ratio — cement parts", 1, N0, "")
inp(18, "ცემენტი : ქვიშა / Cement : sand ratio — sand parts", 4, N0,
    "1:4 ძლიერი, 1:5 ჩვეულებრივი / 1:4 is strong, 1:5 is ordinary")
inp(19, "ცემენტი 1 მ³ ხსნარზე / Cement per m³ of mortar", 350, N0,
    "კგ. 1:4 ≈ 350, 1:5 ≈ 280 / kg. Depends on the mix")
inp(20, "ქვიშა 1 მ³ ხსნარზე / Sand per m³ of mortar", 1.05, N2,
    "მ³ ფხვიერი ქვიშა / m³ of loose sand")
inp(21, "ხსნარის მარაგი / Mortar waste allowance", 0.15, '0%',
    "დაღვრა და ნარჩენი / spillage and what sets in the barrow")
inp(22, "ტომრის წონა / Bag weight", 50, N0, "კგ / kg")

# ---------------- derived ----------------
section(24, "4. გაანგარიშება / THE WORKING")
out(25, "კედლის ფართობი — ბრუტო / Gross wall area", "=B5*B6", N2, "მ² / m²  = perimeter × height")
out(26, "კედლის ფართობი — ნეტო / Net wall area", "=B25-B7", N2, "მ² / m²  after openings")
out(27, "ბლოკი 1 მ²-ზე / Blocks per m²", "=1/((B10+B13)*(B11+B13))", N3,
    "ნაკერის ჩათვლით / each block occupies (L+joint) × (H+joint)")
out(28, "რიგი 1 მ სიმაღლეზე / Courses per metre of height", "=1/(B11+B13)", N3, "")
out(29, "ბეტონის ნაკერი 1 მ²-ზე / Bed-joint mortar per m²", "=B28*1*B13*B12", N3,
    "მ³ / m³  horizontal joints")
out(30, "ვერტიკალური ნაკერი 1 მ²-ზე / Perpend mortar per m²", "=B27*(B11*B12*B13)", N3,
    "მ³ / m³  vertical joints between blocks")
out(31, "ხსნარი 1 მ²-ზე / Mortar per m²", "=B29+B30", N3, "მ³ / m³")

# ---------------- answers ----------------
section(33, "5. საყიდელი / WHAT TO ORDER")
out(34, "ბლოკი — თეორიული / Blocks — bare requirement", "=ROUNDUP(B26*B27,0)", N0,
    "ცალი / pieces, no allowance")
out(35, "ბლოკი — საყიდელი / BLOCKS TO BUY", "=ROUNDUP(B34*(1+B14),0)", N0,
    "ცალი / pieces, waste included", big=True)
out(36, "ხსნარი / Mortar volume", "=B26*B31*(1+B21)", N3, "მ³ / m³, waste included")
out(37, "ცემენტი / Cement", "=B36*B19", N0, "კგ / kg")
out(38, "ცემენტის ტომარა / CEMENT BAGS", "=ROUNDUP(B37/B22,0)", N0,
    "ტომარა / bags", big=True)
out(39, "ქვიშა / SAND", "=ROUNDUP(B36*B20*10,0)/10", N2, "მ³ / m³", big=True)
out(40, "წყალი — მიახლოებით / Water, roughly", "=ROUNDUP(B37*0.5,0)", N0,
    "ლიტრი / litres ≈ half the cement weight")

# ---------------- notes ----------------
section(42, "6. რა არ არის ჩათვლილი / WHAT THIS DOES NOT COVER")
notes = [
    "ფუნდამენტი, არმატურა, ანკერი, ბეტონის სარტყელი — ცალკე ითვლება. / "
    "Foundation, reinforcement, ties and any concrete ring beam are not in this — price them separately.",
    "ღიობები: ჭიშკრის ადგილი გამოაკელით ზემოთ, უჯრა B7. 180 მ პერიმეტრს თითქმის ყოველთვის აქვს შესასვლელი. / "
    "Openings: put the gate area into B7. A 180 m perimeter almost always has a way in.",
    "ბლოკი კიდეზე რომ დაიდოს (10 სმ კედელი), შეცვალეთ B12 — ხსნარი განახევრდება. / "
    "If the blocks go on edge for a 10 cm wall, change B12 — the mortar roughly halves.",
    "ცემენტის რაოდენობა ხსნარის შემადგენლობაზეა დამოკიდებული. B19 შეცვალეთ, თუ ოსტატი სხვა პროპორციას იყენებს. / "
    "Cement per m³ depends on the mix. Change B19 if your mason works to a different proportion.",
    "შედეგი მასალაა, არა ფასი. ფასი ხარჯების ლეჯერში შედის მიწოდების შემდეგ. / "
    "This gives quantities, not cost. The cost goes into the expense ledger when it is bought.",
]
for i, n in enumerate(notes):
    ws.merge_cells(f"A{43+i}:D{43+i}")
    put(f"A{43+i}", "· " + n, NOTE, align="left")
    ws.row_dimensions[43 + i].height = 26

wb.save(OUT)
print("saved", OUT)
